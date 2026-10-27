"""Build the advisor brief, evidence tables, figures, and editable/PDF slides.

Run after audit_data.py and benchmark_text_baseline.py. The curated case manifest
and evidence catalog are reviewed inputs, not automatically inferred labels.
"""
from pathlib import Path
import base64
import csv
import html
import json
import textwrap
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

OUT=Path(__file__).resolve().parent
REPO=OUT.parents[1]
load=lambda p:json.loads(p.read_text())
audit=load(OUT/'dataset_audit.json')
baseline=load(OUT/'text_baseline.json')
papers=load(OUT/'evidence_catalog.json')
cases=load(OUT/'selected_examples.json')
pilot=load(REPO/'experiments/information_gap/results/conclusions.json')
E=html.escape
NAVY='#172b46';TEAL='#117d83';ORANGE='#bd5b36';MUTED='#52647a'

def table(headers,rows):
    return '<div class="scroll"><table><thead><tr>'+''.join('<th>'+E(str(x))+'</th>' for x in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+str(x)+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table></div>'

def data_uri(path):
    p=OUT/path;typ='image/jpeg' if p.suffix.lower() in ['.jpg','.jpeg'] else 'image/png'
    return 'data:'+typ+';base64,'+base64.b64encode(p.read_bytes()).decode()

def fig(path,caption):return '<figure><img src="'+data_uri(path)+'"><figcaption>'+caption+'</figcaption></figure>'

def make_figures():
    plt.rcParams.update({'font.size':12,'axes.spines.top':False,'axes.spines.right':False,'axes.titleweight':'bold','font.family':'DejaVu Sans'})
    fig_,ax=plt.subplots(figsize=(10.8,4.8),layout='constrained')
    keys=['MIMIC_CXR_training','MIMIC_CXR_test','IU_Xray_training','IU_Xray_test']
    vals=[audit['datasets'][k]['counts']['explicit_comparison_reference']/audit['datasets'][k]['n_records']*100 for k in keys]
    labels=['MIMIC local train','MIMIC local test','IU local train','IU local test']
    ax.barh(labels[::-1],vals[::-1],color=[TEAL,TEAL,ORANGE,ORANGE]);ax.set_xlim(0,58)
    for i,(k,v) in enumerate(zip(keys,vals)):
        r=audit['datasets'][k];ax.text(v+.8,3-i,f'{v:.1f}%  ({r["counts"]["explicit_comparison_reference"]}/{r["n_records"]})',va='center',fontsize=11)
    ax.set_xlabel('Records with an explicit comparison reference (%)');ax.set_title('Targets refer to prior studies in single-image tasks',loc='left',pad=18)
    ax.text(0,-.24,'Lexical screen of local task files; records are not independent patients.',transform=ax.transAxes,color=MUTED,fontsize=10)
    fig_.savefig(OUT/'assets/report_comparisons.png',dpi=190);plt.close(fig_)
    fig_,ax=plt.subplots(figsize=(10.8,4.8),layout='constrained')
    models=baseline['SLAKE']['models'];names=['global','type','question_only'];labels=['Global answer mode','Type answer mode*','Question lookup']
    means=np.array([models[k]['accuracy']*100 for k in names]);ci=np.array([models[k]['image_cluster_bootstrap95'] for k in names])*100
    bars=ax.bar(labels,means,color=[MUTED,TEAL,ORANGE],width=.58)
    ax.errorbar(np.arange(3),means,yerr=np.array([means-ci[:,0],ci[:,1]-means]),fmt='none',color=NAVY,capsize=6)
    for b,m,k in zip(bars,means,names):ax.text(b.get_x()+b.get_width()/2,m+4,f'{m:.1f}%\n{models[k]["correct"]}/1061',ha='center',fontsize=12)
    ax.set_ylim(0,73);ax.set_ylabel('Normalized exact-match accuracy (%)');ax.set_title('SLAKE: train-fitted baselines without images',loc='left',pad=18)
    ax.text(0,-.22,'English test: 1061 QA / 96 images; 0 image overlap with train.\n95% intervals bootstrap test images. *Uses content-type metadata.',transform=ax.transAxes,fontsize=10,color=MUTED)
    fig_.savefig(OUT/'assets/slake_text_baseline.png',dpi=190);plt.close(fig_)
    fig_,ax=plt.subplots(figsize=(10.8,4.9),layout='constrained')
    x=np.arange(3);width=.22
    vals=[[10,12,12],[12,12,2],[12,12,11]]
    for i,(v,label,color) in enumerate(zip(vals,['Baseline','Provenance prompt','Evidence-first prompt'],[MUTED,ORANGE,TEAL])):
        bs=ax.bar(x+(i-1)*width,v,width,label=label,color=color)
        for b,n in zip(bs,v):ax.text(b.get_x()+width/2,n+.15,str(n),ha='center',fontsize=11)
    ax.set_xticks(x,['Missing fact:\ncorrectly unknown','Positive fact:\ncorrectly yes','Negative fact:\ncorrectly no']);ax.set_ylim(0,15.5);ax.set_ylabel('Correct choices out of 12');ax.legend(loc='upper center',ncol=3,frameon=False,fontsize=10)
    ax.set_title('Existing pilot: abstention can improve while updating fails',loc='left',pad=16)
    ax.text(0,-.22,'Archived Huatuo-labeled checkpoint results; 12 selected facts.\nOption scores average all six answer-order permutations; no fresh inference.',transform=ax.transAxes,fontsize=10,color=MUTED)
    fig_.savefig(OUT/'assets/prompt_tradeoff.png',dpi=190);plt.close(fig_)

def write_tables():
    with (OUT/'paper_evidence.csv').open('w') as f:
        w=csv.DictWriter(f,fieldnames=list(papers[0]));w.writeheader();w.writerows(papers)
    cols=['case_id','dataset','id','title','evidence_label','question','answer','caption','assessment','aligned_target','proposed_test','source','article_url','article_locator','review_status']
    with (OUT/'selected_examples.csv').open('w') as f:
        w=csv.DictWriter(f,fieldnames=cols,extrasaction='ignore');w.writeheader();w.writerows(cases)

def make_html():
    css='''body{margin:0;background:#f3f6f9;color:#172b46;font:16px/1.65 system-ui,sans-serif}main{max-width:1160px;margin:auto;padding:50px 35px 90px}h1{font-size:42px;line-height:1.16;max-width:1000px}h2{font-size:28px;margin-top:55px;border-top:2px solid #d9e3eb;padding-top:28px}h3{font-size:21px}a{color:#116d80}p{max-width:1000px}.lead{font-size:22px;line-height:1.5}.meta,.small{font-size:13px;color:#52647a}.callout{padding:22px 26px;background:#e4f1f1;border-left:5px solid #117d83;border-radius:4px}.warn{background:#fff0e6;border-color:#bd5b36}table{border-collapse:collapse;width:100%;font-size:14px;background:white}td,th{padding:13px 15px;border:1px solid #dbe4eb;text-align:left;vertical-align:top}th{background:#172b46;color:white}.scroll{overflow-x:auto}figure{margin:28px 0;padding:18px;background:white;border-radius:8px}figure img{max-width:100%;max-height:550px;display:block;margin:auto}figcaption{font-size:13px;color:#52647a;margin-top:15px}.case{background:white;margin:22px 0;padding:22px 28px;border-radius:8px}.case img{max-width:100%;max-height:440px}.grid{display:grid;grid-template-columns:1fr 1fr;gap:26px}.tag{font-size:12px;background:#e4f1f1;padding:4px 8px;border-radius:5px}.formula{font:18px/1.9 monospace;background:#e7edf3;padding:20px;overflow:auto}summary{cursor:pointer;color:#116d80}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#eef3f7;padding:14px;font-size:12px}input{width:95%;padding:12px;border:1px solid #c2d1df;font-size:16px}nav{display:flex;flex-wrap:wrap;gap:18px}li{margin:8px 0}@media(max-width:750px){main{padding:22px}.grid{display:block}h1{font-size:32px}}@media print{body{background:white}main{padding:0}nav,input{display:none}h2{break-before:page}table{font-size:10px}.case{break-inside:avoid}details{display:block}}'''
    s=['<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Medical MLLM information gap — advisor evidence brief</title><style>'+css+'</style><main>']
    s+=['<div class="meta">RESEARCH ASSESSMENT · Literature cutoff: October 19, 2026 · local audit and source inspection</div><h1>Which medical claims are justified by the evidence the model actually receives?</h1>',
    '<p class="lead">The information-gap framework is reasonable and worth discussing with your advisor. The evidence supports a concrete supervision problem and a testable research program. It does not yet establish a dominant root cause of medical MLLM failure.</p>',
    '<nav><a href="#judgment">Judgment</a><a href="#papers">Paper evidence</a><a href="#audit">Local audit</a><a href="#cases">Data examples</a><a href="#pilot">Existing pilots</a><a href="#theory">Theory</a><a href="#study">Decisive study</a><a href="#reproduce">Files and reproduction</a></nav>',
    '<p><a href="advisor_slides.pdf">Presentation PDF</a> · <a href="advisor_slides.pptx">Editable PowerPoint</a> · <a href="paper_evidence.csv">Evidence CSV</a> · <a href="selected_examples.csv">Case CSV</a> · <a href="selected_examples.jsonl">Full case manifest</a></p>',
    '<div class="callout"><b>Suggested thesis.</b> Some medical training targets combine visible observations with facts supplied only to the annotator or synthetic teacher. Explicitly matching each claim to its available evidence may reduce unsupported assertions while retaining useful medical knowledge and diagnostic performance.</div>']
    s+=['<h2 id="judgment">Assessment of the discussion notes</h2>',table(['Notes claim','Judgment','Revision for the advisor'],[
    ['Bidirectional image–text gap','Agree as an organizing framework','Separate unavailable patient facts, omitted visible findings, ambiguous findings, and available-but-unused evidence. They need different labels and experiments.'],
    ['More data mainly improves language; scaling does little for evidence','Too strong','Lingshu and Huatuo report gains; MedGemma 1.5 expands evidence access. Test whether scaling <em>alone</em> leaves specific unsupported-claim rates unresolved, at matched data quality. [P02–P04]'],
    ['MLE on image-only reports encourages unsupported content','Plausible; concrete cases exist','A Bayes-optimal conditional distribution can marginalize missing context honestly. MLE does not mathematically force fabrication; deterministic reporting and answer-only evaluation can incentivize over-specific assertions.'],
    ['Unreported content gets no gradient','Literal wording is wrong','There is no direct target naming that content, but shared visual features, other captions and previous training supply indirect learning. A coverage deficit needs independent image labels.'],
    ['Teacher captions are mostly off-image; mixing necessarily shrinks overlap','Not measured','Source rankings in the notes are qualitative hypotheses. Mixture entropy and usable-information ratios are not monotonic in source size; answerability must be audited per source.'],
    ['Frozen representations bound downstream recovery','Agree with conditions','The limitation is about a fixed input/representation, not the calendar stage called post-training. Encoder updates, crops, additional views and tools change the information available.'],
    ['A failed probe means pretraining is needed','Disagree as a decision rule','Probe failure may reflect sample size, pooling or predictor capacity. Successful decoding also does not prove causal use of a lesion. Interventions are needed.'],
    ['Counterfactual tests and provenance are preferable to hedging rate','Agree; strong direction','Require appropriate changes after relevant evidence changes, and stability after irrelevant changes. Preserve performance on answerable cases. Very close methods already exist. [P15,P19–P22]'],
    ['No current signal rewards grounding or uncertainty','Incorrect as a universal claim','MCR, RadVLM, RULE, MedReasoner and CARE-MVLM already address aspects of these objectives. The remaining claim must be narrower.'],
    ['Medical dense recaptioning cannot transfer','Too absolute','RoI- and mask-based medical supervision already exists. Its cost, coverage and independent reliability remain constraints, rather than a proof of impossibility.'],
    ['Typed uncertainty follows automatically from PVI','Not established','Ablation scores measure model/intervention dependence. They do not label why evidence is insufficient; experts or structural evidence labels are still needed.'],
    ['All three failures share one demonstrated cause','Not established','Missing context is a shared candidate mechanism. Resolution, acquisition, disease ambiguity, prevalence, label noise and deployment shift remain competing or interacting causes.']])]
    s+=['<h3>Repair the interpretation of the notes’ cited evidence</h3>',table(['Citation in notes','What it can support','What to avoid'],[
    ['[1] MIRAGE via the 3D oncology paper','Cite MIRAGE directly [P13]; the oncology preprint is a useful separate private-data study.','Conflating private data with guaranteed visual dependence; text cues can remain. <a href="https://pubmed.ncbi.nlm.nih.gov/42818963/">Oncology source</a>.'],
    ['[2] RadImageNet-VQA','Text-only performance depends on benchmark and question format [P14].','Calling 2–10% open-ended accuracy a mathematically defined chance level.'],
    ['[3] ED reporting study','Clarity and diagnostic quality can diverge [P05].','Treating 1434 reader ratings as independent patients or attributing errors to scaling objectives. AIRead is also a counterexample to a blanket failure claim.'],
    ['[4] Lingshu physician agreement','36% exact rating agreement, 81% within one point, weighted kappa 0.64 [P03].','Interpreting exact ordinal-score disagreement as evidence that 64% of diagnostic ground truths are unreliable.'],
    ['[5] Wound study','A small study raises actionability concerns [P25].','Calling hedging miscalibrated without labels for what uncertainty is justified; the scoring rule penalizes hedging by construction.'],
    ['[6] MedBench v5','Dynamic evidence stress testing is relevant [P24].','Presenting 720 derived scenarios as independent cases or assuming all process judges have clinical validation.'],
    ['[7] Paediatric imaging review','Use any reported misses as illustrative anecdotes after checking the underlying case.','Using a review anecdote as a controlled demonstration of context dependence or as prevalence evidence.'],
    ['[8] Non-existent priors','A particularly strong direct precedent [P06].','Presenting missing-prior cleaning or DPO suppression [P07] as new.'],
    ['Our token-zeroing claim','A hypothesis for a targeted lesion intervention.','I did not locate and independently validate the cited token-zeroing experiment in this review. Zeroing can introduce artifacts or target irrelevant tokens.']])]
    s+=['<h2 id="papers">Evidence from backbones and medical reasoning papers</h2><p>Rows distinguish an author-reported result from our interpretation. Follow the linked primary source and locator for the exact protocol. This is a targeted review, not an exhaustive systematic review. Scores from different papers are not directly comparable.</p><input id="papersearch" placeholder="Filter papers, methods, or limitations…" aria-label="Filter paper evidence">']
    for p in papers:
        s.append('<article class="paper case" id="'+p['id']+'"><span class="tag">'+p['id']+' · '+E(p['relation'])+'</span><h3><a href="'+E(p['url'])+'">'+E(p['paper'])+'</a></h3><p class="meta">'+E(p['year'])+' · Read: '+E(p['locator'])+'</p><p><b>Reported evidence:</b> '+E(p['finding'])+'</p><p><b>Interpretation:</b> '+E(p['implication'])+'</p><p><b>Boundary:</b> '+E(p['limit'])+'</p></article>')
    s+=['<h3>Where a new contribution could remain</h3><p>MIRAGE and Beyond Accuracy already challenge image dependence. MCR and Med-OPD already use image-conditioned reward signals. CARE-MVLM already uses lesion removal, matched random removal and abstention. MedTRACE already models evidence sufficiency and acquisition. A defensible candidate contribution is a <b>claim-level audit and controlled intervention on teacher-only patient facts and missing panels</b>, evaluated on correct-answer/false-premise cases as well as incorrect diagnoses. Claim novelty only after a broader overlap check.</p>',
    '<h2 id="audit">What the local data actually show</h2><p>The scripts scan local files read-only and save checksums. Counts describe these copies and task wrappers, not the complete published datasets. Multiple records can belong to one patient or figure. Keyword matches are screening flags, not clinical error labels.</p>']
    rows=[]
    for k in ['MIMIC_CXR_training','MIMIC_CXR_test','IU_Xray_training','IU_Xray_test']:
        r=audit['datasets'][k];n=r['n_records'];c=r['counts']['explicit_comparison_reference'];b=r['counts']['longitudinal']
        rows.append([E(k),f'{n:,}',f'{c:,} / {n:,} ({c/n:.1%})',f'{b:,} / {n:,} ({b/n:.1%})'])
    s.append(table(['Local task file','Records; all have one image','Explicit comparison references','Broader temporal screen'],rows))
    s+=['<p>The explicit screen looks for phrases such as “compared with,” “in comparison with,” and “previous study.” The broader screen also includes “stable,” “unchanged” and follow-up wording. Neither count asserts that every matched sentence is false. In these wrappers, no prior image, report or history is provided in the prompt; inspect the exact target before editing.</p>',fig('assets/report_comparisons.png','New audit. Denominators are local records; no patient-level confidence interval is claimed.'),
    '<h3>PubMedVision: source information can survive reformatting</h3>']
    r=audit['datasets']['PubMedVision_instruction'];n=r['n_records'];c=r['numeric_age_provenance']
    s.append(table(['Screen','Count','Interpretation'],[
    ['Instruction records in local file',f'{n:,}','Published dataset card lists 647,031; this local copy has 272 fewer.'],
    ['Local original-caption records',f'{audit["datasets"]["PubMedVision_original_captions"]["n_records"]:,}','Do not assume complete one-to-one caption coverage.'],
    ['Answer contains numeric age',str(c['answer_numeric_age']),'Lexical pattern only.'],
    ['An answer age is absent from the question',str(c['answer_age_absent_question']),'Exact numeric set difference; does not inspect image text.'],
    ['That age also appears in the joined source caption',f'{c["source_shared_age_absent_question"]} / {n:,} ({c["source_shared_age_absent_question"]/n:.3%})','A provenance-screen frequency in this file, not a validated mismatch rate.'],
    ['Answer mentions an external-test keyword',f'{r["answer_screen_counts"]["external_test_event"]:,}','Includes actual histology images, hypothetical recommendations and legitimate uncertainty; never equate with missing evidence.']]))
    s+=['<p>Image OCR, the identity of the patient referenced, hypothetical examples and approximate-age questions are not adjudicated by the age count. The 993 count is therefore neither a clinical hallucination rate nor a lower bound on one. A uniform 4,000-record reservoir is saved for independent annotation; keyword-selected examples are illustrative.</p>',
    '<h3>Benchmark controls fitted without images</h3>']
    b=baseline['SLAKE'];rows=[]
    for key,title in [('global','Global train answer mode'),('type','Train mode by content-type metadata'),('question_only','Exact normalized train-question lookup; global fallback')]:
        v=b['models'][key];lo,hi=v['image_cluster_bootstrap95'];rows.append([title,f'{v["correct"]}/{v["n"]} = {v["accuracy"]:.1%}',f'{lo:.1%}–{hi:.1%}'])
    s+= [table(['SLAKE English baseline','Test exact match','95% image-cluster bootstrap interval'],rows),
    '<p>Train: 4,919 English QA over 450 images. Test: 1,061 QA over 96 images. Image overlap is zero. Of test questions, 999 have an exact normalized question in train; answers are selected using train frequencies only. Question lookup gets 494/913 on the designated visual-QA subset and 87/148 on knowledge QA under the metadata-fallback variant. Repeated templates and answer frequency explain useful baseline performance without proving any particular image claim unsupported. These are new lookup results, not medical MLLM scores.</p>',
    fig('assets/slake_text_baseline.png','Question lookup uses no image, type metadata or pretrained LM. Type mode uses content-type metadata. Lowercasing, punctuation and whitespace normalization are fixed.'),
    '<p>On PMC-VQA, the most frequent training answer letter is B. Predicting B blindly gives 15,854/50,000 = <b>31.708%</b>, compared with a nominal four-option random baseline of 25%. The local files contain 14 malformed train labels and 10 malformed test labels; this raw-file denominator includes them. Option balancing and shuffling are necessary controls, but this result does not establish image independence for individual questions.</p>',
    '<p>The local VQA-RAD file named MED_RAD_test.json contains 2,248 rows over 314 images, resembling the full original corpus rather than a verified held-out split. Its CLOSED metadata also includes non-binary answers. Do not infer published test scores or assume all CLOSED questions are yes/no. The LLaVA-Med code checkout exists, but the expected data/instruct and data/pmc_articles training directories are empty; this review makes no local training-example claim for that dataset.</p>',
    '<h2 id="cases">Eleven traceable examples, including controls</h2><p>Full original QA, captions, file paths and intervention suggestions are in the manifest. Labels below concern structural evidence availability; subtle image findings and diagnostic appropriateness still need specialist adjudication. No new patient-specific medical advice is implied.</p>']
    for r in cases:
        s.append('<article class="case" id="'+r['case_id']+'"><span class="tag">'+r['case_id']+' · '+E(r['dataset'])+' · '+E(r['evidence_label'])+'</span><h3>'+E(r['title'])+'</h3><p class="meta">Record '+E(r['id'])+' · '+E(r.get('article_locator','local record'))+'</p>')
        if r.get('assets'):
            if r.get('original_figure_asset'):
                s.append('<div class="grid">'+fig(r['assets'][0],'Actual student image: final-timepoint column only.')+fig(r['original_figure_asset'],'Original Figure 3: MRI 1, MRI 2 and MRI 3. Earlier columns and headings are absent from the student image.')+'</div>')
            else:s.append(fig(r['assets'][0],'Actual local image supplied with this QA; displayed without generated alterations.'))
        s.append('<p><b>Question:</b> '+E(r['question'])+'</p><p><b>Assessment:</b> '+E(r['assessment'])+'</p><p><b>Aligned target behavior:</b> '+E(r['aligned_target'])+'</p><p><b>Discriminating test:</b> '+E(r['proposed_test'])+'</p>')
        if r.get('article_url'):s.append('<p><a href="'+E(r['article_url'])+'">Original article / figure attribution</a></p>')
        s.append('<details><summary>Inspect exact target, caption, and local provenance</summary><p><b>Original answer:</b> '+E(r['answer'])+'</p><p><b>Original caption:</b> '+E(r.get('caption','Not supplied for this task wrapper.'))+'</p><pre>'+E(json.dumps({k:r.get(k) for k in ['source','id','split','row_index','image_metadata','image_relative','review_status'] if r.get(k) is not None},indent=2))+'</pre></details></article>')
    s+=['<h2 id="pilot">What the existing workspace pilots add</h2><p>These runs predate this review. I inspected archived raw-choice files, result summaries, checkpoints and protocols; I did not rerun GPU inference. They are exploratory, use selected cases, and are not independent evidence of population prevalence.</p>',
    '<p>The 96-record evidence-mismatch pilot uses 32 examples in each of external-context, diagnostic-inference and visual-description strata. Both tested checkpoints show lower image-identity gain and greater caption specificity for the external-context targets. This is consistent with model dependence on source information. Differences in claim content and selection remain confounds, and caption rescue can partly expose the answer directly.</p>',
    '<p>The follow-up uses 12 reviewed external facts, plus 12 broad-anatomy images (six head/chest pairs). A provenance prompt increases missing-fact abstention from 10/12 to 12/12 for the Huatuo-labeled checkpoint, but correct use of explicitly negative context falls from 12/12 to 2/12. For MedSIGHT, the same prompt improves missing-fact abstention from 6/12 to 12/12 but reduces correct anatomy under conflicting notes from 12/12 to 1/12. These are constrained choice-score results, not free-text clinical error rates.</p>',
    fig('assets/prompt_tradeoff.png','From experiments/information_gap/results/conclusions.json; all six answer-order permutations are averaged. The small selected sample makes these diagnostic examples, not a model ranking.'),
    '<p>Coarse head/chest anatomy is decodable and used in both checkpoints, including baseline decisions under conflicting notes. That is a useful negative result against a blanket coverage-gap story. It says nothing decisive about small lesions or unreported findings. The Huatuo-labeled local checkpoint has a LLaVA-Qwen2 architecture; do not silently equate it with the original 2024 HuatuoGPT-Vision-34B paper configuration.</p>',
    '<h2 id="theory">Theory worth keeping, with narrower claims</h2>',
    '<p>Let Y be a finite finding label, I the complete supplied image input, C omitted context, and Z=g(I) a deterministic frozen representation. Hold the question fixed, or condition every term below on the question. With log base 2 and a common population distribution:</p>',
    '<div class="formula">E[−log₂ q(Y|Z)]<br>= H(Y|I,C) + I(Y;C|I) + I(Y;I|Z)<br>  + E<sub>Z</sub> KL₂(p(Y|Z) || q(Y|Z)).</div>',
    '<p>The identity follows from cross-entropy = conditional entropy + KL, then H(Y|Z) = H(Y|I) + I(Y;I|Z), then H(Y|I) = H(Y|I,C) + I(Y;C|I). Deterministic Z ensures the required conditional independence. If context is actually supplied, define the representation and conditioning accordingly. The decomposition is a population identity, not a per-record diagnosis of the model’s error.</p>',
    '<p>The last term includes approximation, fitting and calibration error, rather than only a failure to look at retained evidence. Lost encoding is not identical to absent report coverage: input resizing, frame omission and representation compression can lose information about reported findings too. Reader disagreement belongs in H(Y|I,C) only if Y is the reader label or its noise affects the chosen target; it is not automatically biological irreducibility.</p>',
    '<div class="formula">P(error) ≥ max{0, [H₂(Y|Z) − 1] / log₂ K}, K = |Y| ≥ 2.</div>',
    '<p>This coarse Fano bound is valid for predictors restricted to frozen Z. It is often vacuous, especially for binary labels, and gives no useful numerical floor without estimating H(Y|Z). It does not prove a universal pretraining/post-training boundary. Updating the encoder or acquiring evidence changes the premise.</p>',
    '<p><b>Ablation ≠ information estimation.</b> A single model’s log q(y|I) − log q(y|no image) measures that model under a chosen intervention, which may be out of distribution. It can be negative even when true image information is positive. The archived synthetic theory check exhibits exactly that. Proper PVI uses specified predictor families, comparable fitted null/input predictors and held-out evaluation [T01]. Neither PVI nor a probe automatically assigns an expert evidence-sufficiency label.</p>',
    '<p><b>Conformal calibration is a later layer.</b> Choose the actual loss, calibration units and monotone/nested output family before invoking a theorem [T02–T03]. Standard expected-risk control is not per-patient certainty or automatically a guarantee among non-abstained cases. Clinical labels must distinguish missed findings, false findings and unsupported attribution; reports alone are insufficient. Hold calibration patients aside from score design and test external sites separately. Rare-finding or site-specific guarantees can require far more than a few hundred cases.</p>',
    '<h2 id="study">The decisive study to propose</h2>',
    '<p>Start with <b>one modality and two evidence requirements</b>: current-versus-prior CXR comparisons, and image-versus-pathology claims in PMC-derived VQA. Do not begin with broad pretraining, a learned information decomposition and conformal deployment guarantees simultaneously.</p>']
    s.append(table(['Stage','Design','What it can decide'],[
    ['1. Independent audit','Uniform samples across sources, not only keyword hits. Two qualified readers annotate atomic claims: visible, contextual, derived, unavailable, ambiguous; mark supporting view/timepoint and disagreement. Blind readers to model identity.','Estimate claim- and example-level mismatch separately, with patient/article clustering and adjudicated intervals. Include visibly unreported findings to test the coverage hypothesis.'],
    ['2. Factorial evidence test','Hold the question and current image fixed. Add/withhold the actual prior or pathology report; separately change irrelevant context. Use matched original image pairs and only validated lesion edits.','Missing-evidence attribution versus legitimate context updating, rather than generic refusal. Supply provenance and authority so contradiction tests are well defined.'],
    ['3. Matched training intervention','Same records, tokens, steps, backbone and initialization: original targets; cleaned/provenance targets; original targets plus MCR-like objective; aligned targets plus that objective. Add a size-matched random-removal control.','A causal effect of supervision alignment beyond data quantity, answer frequency or generic grounding rewards.'],
    ['4. Representation intervention','Independently labeled lesion-present/absent data; matched acquisition/anatomy. Compare held-out probes at encoder/projector/LLM layers and actual decisions. Then contrast frozen versus updated encoder with controlled budgets.','Locate a practical bottleneck. A failed finite probe alone cannot establish lost information, and a successful probe alone cannot establish causal lesion use.'],
    ['5. Generalization','Split by patient and source article before generating QA. Freeze evaluation prompts, masks, annotation policy and model checkpoints. Reserve distinct calibration and external-test patients.','Avoid near-duplicate figures, reference leakage, template memorization and tuning on the evaluation cases.']]))
    s+=['<h3>Primary outcome: joint evidential success</h3><p>Count a case successful only when the model answers adequately with sufficient evidence, preserves the answer under an irrelevant intervention, removes or qualifies the affected claim when its evidence is withheld, and updates correctly when valid evidence arrives. Report these components separately as well. For ambiguous diagnoses use probabilistic targets or expert acceptable sets, rather than demanding a unique definitive label.</p>',
    '<p>Measure unsupported claim rate, current-finding sensitivity, false-positive rate, relevant/irrelevant intervention response, evidence-acquisition choice, and risk versus coverage. Include the rate of <b>correct final answer with a false or unavailable visual premise</b>. Withholding an image needed for a question should not be scored identically to obtaining a negative image that supports “absent.” A no-lesion answer and insufficient evidence are different outputs.</p>',
    '<h3>Falsification and practical decision rules</h3>',
    '<ul><li>If independent readers judge most sampled targets adequately supported, the broad mismatch-prevalence hypothesis weakens.</li><li>If aligned-target training does not improve held-out joint success beyond equal-size random cleaning or MCR-style training, the causal training claim weakens.</li><li>If gains come entirely from extra abstention with fewer correct answerable cases, the proposed mitigation fails.</li><li>If added valid context does not restore factual specificity, investigate context use and the uncertainty policy before adding pretraining.</li><li>If a stronger frozen readout or improved crops resolve the finding, do not attribute failure to irrevocably lost encoder information.</li></ul>',
    '<div class="callout"><b>Advisor decision.</b> This is ready as a scoped research hypothesis with traceable evidence. Ask the advisor to choose a modality, secure annotation expertise, and agree on the matched supervision intervention and primary joint metric. Present the coverage mechanism and the dominant-root-cause claim as open questions.</div>',
    '<h2 id="reproduce">Files and reproduction</h2><p>New artifacts are confined to this folder. The source datasets and pre-existing experiment outputs were not modified. The report embeds displayed images and plots for offline reading; source links require internet. The slides contain source IDs and speaker notes. Dataset use remains subject to each original source’s terms.</p>',
    '<pre>python -m pip install ijson numpy matplotlib pillow python-pptx pymupdf\npython Plans/medical_mllm_gap_review/audit_data.py\npython Plans/medical_mllm_gap_review/benchmark_text_baseline.py\npython Plans/medical_mllm_gap_review/build_review.py</pre>',
    '<p>The selected case manifest and evidence catalog are curated inputs; rebuilding does not replace specialist adjudication. The audit stores a reproducible uniform reservoir and a separate set of screening examples. No model inference is performed by these commands.</p>',
    '<ul><li><a href="dataset_audit.json">Dataset counts, regexes and input SHA256 checksums</a></li><li><a href="text_baseline.json">Baseline definitions, counts and intervals</a> · <a href="text_baseline_predictions.csv">Per-question predictions</a></li><li><a href="pubmedvision_random_sample.jsonl">Uniform PubMedVision reservoir, 4,000 records</a></li><li><a href="selected_examples.jsonl">Full curated examples with provenance</a></li><li><a href="evidence_catalog.json">Machine-readable paper evidence</a></li><li><a href="provenance.json">Review scope and source-note checksum</a></li><li><a href="../../experiments/evidence_mismatch/results/report.html">Existing mismatch pilot</a> · <a href="../../experiments/information_gap/results/report.html">Existing information-gap pilot</a></li></ul>',
    '<p class="small">The existing pilot links above are available within the repository. Evidence strength: structural input availability is directly auditable; fine diagnostic sufficiency and causal training effects remain to be established.</p></main><script>document.getElementById("papersearch").addEventListener("input",function(){let q=this.value.toLowerCase();document.querySelectorAll(".paper").forEach(e=>e.hidden=!e.textContent.toLowerCase().includes(q))})</script></html>']
    # This folder is one level below Plans; existing experiments are two levels above.
    (OUT/'advisor_evidence_brief.html').write_text('\n'.join(s))

class Slides:
    def __init__(self):
        self.prs=Presentation();self.prs.slide_width=Inches(13.333333);self.prs.slide_height=Inches(7.5)
        from matplotlib.font_manager import findfont, FontProperties
        self.pdf=pymupdf.open();self.info=[];self.font=findfont(FontProperties(family='DejaVu Sans'))
        self.bfont=findfont(FontProperties(family='DejaVu Sans',weight='bold'))
    def rect(self,x,y,w,h,color):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color[1:]);shape.line.fill.background()
        rgb=tuple(int(color[i:i+2],16)/255 for i in (1,3,5));self.page.draw_rect(pymupdf.Rect(x,y,x+w,y+h),fill=rgb,color=None)
    def text(self,txt,x,y,w,h,size=19,color=NAVY,bold=False):
        rgb=tuple(int(color[i:i+2],16)/255 for i in (1,3,5));font='B' if bold else 'R'
        self.page.insert_font(fontname=font,fontfile=self.bfont if bold else self.font)
        # Use the same final point size in both exports. Shrink only when needed.
        original=size
        while True:
            sh=self.page.new_shape();ret=sh.insert_textbox(pymupdf.Rect(x,y,x+w,y+h),txt,fontname=font,fontsize=size,color=rgb,lineheight=1.22)
            if ret>=0 or size<=10:sh.commit();break
            size-=.5
        if ret<0:raise ValueError('Slide text overflow: '+txt[:50])
        if size<original-2:self.info[-1].setdefault('shrunk_text',[]).append({'text':txt[:70],'from':original,'to':size})
        box=self.slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72));tf=box.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
        for i,line in enumerate(txt.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:]);p.space_after=Pt(2);p.line_spacing=1.12
    def pic(self,path,x,y,w,h):
        path=OUT/path;im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale;xx=x+(w-nw)/2;yy=y+(h-nh)/2
        self.slide.shapes.add_picture(str(path),Inches(xx/72),Inches(yy/72),width=Inches(nw/72),height=Inches(nh/72));self.page.insert_image(pymupdf.Rect(xx,yy,xx+nw,yy+nh),filename=str(path))
    def new(self,title,refs='',notes=''):
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6]);self.page=self.pdf.new_page(width=960,height=540)
        self.info.append({'title':title,'refs':refs,'notes':notes})
        self.rect(0,0,960,540,'#f6f8fb');self.rect(0,0,960,8,TEAL)
        self.text('MEDICAL MLLMs  /  EVIDENCE & SUPERVISION',40,25,850,20,10,MUTED,True)
        self.text(title,40,59,880,79,28,NAVY,True)
        self.rect(40,496,880,1,'#d4dee7')
        self.text(refs,40,505,835,26,9,MUTED)
        self.text(str(len(self.info)),890,504,30,18,10,MUTED)
        self.slide.notes_slide.notes_text_frame.text=notes+'\n\nSources: '+refs
    def bullets(self,items,x=45,y=150,w=865,h=320,size=21):
        step=h/len(items)
        for i,t in enumerate(items):
            self.rect(x,y+i*step+8,5,5,TEAL);self.text(t,x+21,y+i*step,w-21,step-8,size)
    def save(self):
        self.prs.save(OUT/'advisor_slides.pptx');self.pdf.save(OUT/'advisor_slides.pdf');self.pdf.close();(OUT/'slides_manifest.json').write_text(json.dumps(self.info,indent=2))

def make_slides():
    d=Slides()
    d.new('Which claims can the supplied evidence justify?', 'Advisor discussion · October 19, 2026', 'Assessment of med_mllm_gap_notes.pdf. The framework is reasonable; causal dominance, population prevalence, and the pre/post boundary are hypotheses. New work: read-only data audit, lookup baselines and source-figure inspection. Existing GPU runs are inspected, not rerun.')
    d.text('A tractable research question',45,155,870,38,22,TEAL,True)
    d.text('Does aligning each training claim with the evidence actually supplied reduce unsupported assertions without losing useful diagnostic performance?',45,210,845,140,30,NAVY,True)
    d.text('Central idea: supported. Dominant root cause: unproven.\nCandidate contribution: source provenance + matched supervision intervention.',45,385,860,70,19)

    d.new('Separate four problems before choosing a method','Discussion notes, pp. 2–5', 'Medical knowledge remains useful. Missing evidence, missed perception, ambiguity and unused evidence are not interchangeable. A confirmed diagnosis may be a valid probabilistic target while a fabricated visual premise remains defective.')
    for i,(title,desc) in enumerate([('Missing evidence','A biopsy result or prior study is absent.'),('Missed perception','The relevant finding is present but missed.'),('Ambiguous evidence','The supplied image cannot resolve the diagnosis.'),('Unused evidence','A retained finding loses to a text shortcut.')]):
        x=45+(i%2)*445;y=150+(i//2)*160;d.rect(x,y,425,140,'#e8eff4');d.text(title,x+18,y+18,389,35,22,TEAL,True);d.text(desc,x+18,y+67,389,58,19)

    d.new('Teacher knowledge can exceed student evidence','P01 LLaVA-Med · P02 HuatuoGPT-Vision · P03 Lingshu', 'LLaVA-Med uses caption and inline article mentions with a text teacher. Huatuo uses an image-aware GPT-4V reformatter, so it must not be called caption-only. Lingshu uses image, RoI, metadata and knowledge; reasoning distillation supplies the known answer. Privileged supervision is useful when targets remain answerable or probabilistic; per-claim auditing is needed.')
    for i,(t,b) in enumerate([('Article / teacher','Images + captions\nHistory, pathology,\nother panels / dates'),('Training record','Selected image(s)\n+ generated question\n+ generated answer'),('Possible mismatch','Answer contains a fact\nwhose evidence never\nreaches the student')]):
        x=45+i*300;d.rect(x,180,270,205,'#e5eff1');d.text(t,x+18,200,235,40,21,TEAL,True);d.text(b,x+18,260,235,112,20)
    d.text('This is a risk to measure, not proof that all synthetic supervision is defective.',45,425,865,45,19)

    d.new('Accuracy can rise without stronger visual dependence','P08 MedVLM-R1 · P11 MedVLThinker · P12 Beyond Accuracy', 'Scores are protocol-specific and not cross-paper rankings. MedVLM-R1 reports substantial RL gains. MedVLThinker shows stronger gains from curated text-only training in its comparisons; data quality is a confound. Beyond Accuracy is a small preprint using 100 items per benchmark. Its image-invariance heuristic is not an expert hallucination label; a shuffled image can preserve the correct answer.')
    d.bullets(['MedVLM-R1: 55.11% → 78.22% in its evaluation, using 600 training QA.', 'MedVLThinker: curated text-only RL can outperform its tested multimodal training recipe.', 'Beyond Accuracy: accuracy 56.5% → 58.8%, while image sensitivity falls 48.2% → 39.8%.', 'Test answer correctness and evidence dependence separately. Neither alone certifies a visual premise.'],size=20)

    d.new('The latest work narrows the novelty claim','P15 MCR · P19 Med-OPD · P20 MedTRACE · P21 spatial/factual RL · P22 CARE-MVLM', 'MCR already rewards perceptual and reasoning consistency. Med-OPD weights tokens using teacher image-degradation dependence. MedTRACE models sufficiency and acquisition. CARE-MVLM uses lesion and random-region removals with an answerability selector. New research must add evidence beyond these, not relabel them.')
    d.bullets(['MCR and Med-OPD: image-conditioned rewards and token weighting already exist.', 'MedTRACE: evidence sufficiency, tools, verification and confidence already exist.', 'CARE-MVLM: lesion removal + matched random removal + abstention already exist.', 'Candidate extension: teacher-only patient facts and missing panels, tested through a controlled target-alignment intervention.'],size=20)

    d.new('Single-image report targets refer to unavailable priors','New local-file audit · audit_data.py', 'An explicit-comparison lexical screen, not a clinical hallucination rate. All audited report records contain one image and a generic report prompt, with no supplied prior. Local record counts are not full dataset or patient counts. Broader stable/unchanged flags are reported separately in the brief.')
    d.pic('assets/report_comparisons.png',45,140,870,335)

    d.new('A traceable report example','C10–C11 · MIMIC_CXR/data_report/training.json · rows 0–1', 'Row 0 supplies one image and asks for a concise report. The target compares to a previous radiograph and claims a right-hilar atelectasis is less severe. We do not adjudicate the current radiograph or assert that its present findings are wrong. This case concerns missing comparison evidence in the task wrapper. Missing-prior rewriting and DPO suppression have direct precedents P06–P07.')
    d.rect(45,150,870,85,'#e3edef');d.text('Input: one current image + “generate a concise medical report.”',63,174,832,54,23,NAVY,True)
    d.text('Target: “As compared to the previous radiograph, no relevant change has occurred.”',50,270,850,85,25,ORANGE,True)
    d.text('Current findings can remain valid.\nThe comparison needs another study. Supply it or revise the target.',50,380,850,75,22)

    d.new('A stronger case: earlier timepoints were cropped away','C05 · PubMedVision Instruction-Tuning_394494 · Original: PMC7300102, Figure 3', 'The local 215×123 image was visually compared to the published figure and matches its rightmost MRI 3 column. The generated answer discusses improvement versus earlier scans and exact ages. Structural availability is verified; lesion grades are not independently clinically adjudicated. Restore source panels/dates separately from age metadata. Original figure attribution: Horga et al., Skeletal Radiology, 2020; https://pmc.ncbi.nlm.nih.gov/articles/PMC7300102/.')
    d.text('Actual input',45,145,240,30,18,TEAL,True);d.pic('assets/C05_pmc_394494_0.jpg',45,185,220,205)
    d.text('Published figure',310,145,590,30,18,TEAL,True);d.pic('assets/original_knee_figure.jpg',310,185,590,205)
    d.text('Target describes recovery across time. Input contains only the final timepoint.',45,420,870,52,22,ORANGE,True)

    d.new('A true diagnosis can have unsupported provenance','C01 · PubMedVision Instruction-Tuning_422538 · Original: PMC7720671, Figure 5', 'The supplied file is an MR image; question asks about a large mass. The source caption reports a biopsy confirming fibroadenoma with lactational changes, but no biopsy report is supplied to the student. The claim that biopsy confirmation occurred is off-input even if the final diagnosis is true. We do not independently adjudicate diagnostic certainty from this image. Article attribution: Pregnancy-lactation cycle: how to use imaging methods for breast evaluation, Radiologia Brasileira 2020.')
    d.pic('assets/C01_pmc_422538_0.jpg',45,145,335,300)
    d.text('Original target includes:\n“Our biopsy results confirm this diagnosis.”',420,160,490,118,24,ORANGE,True)
    d.text('MRI features ≠ evidence that a biopsy happened.\n\nKeep the differential; add the pathology record before asserting confirmation.',420,300,490,150,21)

    d.new('Local benchmark scores have strong text baselines','New SLAKE English train/test lookup · benchmark_text_baseline.py', 'Exact normalized train-question lookup with global train-answer fallback; no images, metadata or pretrained language model. Correct: 579/1061 = 54.57%; 95% image-cluster bootstrap interval 51.74–57.21%. There are 4919 training QA over 450 images and 1061 test QA over 96 images. Image overlap is zero; 999 test questions occur in train. Type mode is a separate metadata-aware baseline. This does not claim the same performance protocol as published MLLM evaluations.')
    d.pic('assets/slake_text_baseline.png',45,140,870,335)

    d.new('Controls prevent an overbroad “language prior” story','C06–C09 · dataset_audit.json · text_baseline.json', 'Exact age, biopsy, and temporal regex hits need human review. PMC-VQA test row 9 actually supplies an IHC-containing image, though the exact marker identity can still require context. SLAKE includes legitimate knowledge QA. The 31.708% PMC-VQA answer-letter baseline includes 10 malformed test labels in its denominator. Broad anatomy is demonstrably retained and used in the archived pilot; small-lesion coverage remains open.')
    d.bullets(['Some knowledge is intended: asking the function of an organ legitimately needs medical knowledge.', 'A biopsy/IHC keyword may accompany an actual histology image. A keyword hit is not a missing-evidence label.', 'PMC-VQA: predicting the train-mode option B gives 31.708% on test. Check answer-position balance.', 'Visible dimensions can be answerable while longitudinal percentage change is not. Score atomic claims.'],size=20)

    d.new('Existing pilot: abstention can conceal evidence suppression','Archived information_gap pilot · 12 selected facts · no new GPU inference', 'The local Huatuo-labeled checkpoint uses a LLaVA-Qwen2 architecture, not necessarily the original 2024 Huatuo paper checkpoint. Choices are scored across all six answer permutations. The provenance prompt improves abstention when facts are missing but suppresses negative-context updating. The MedSIGHT counterpart also suppresses visual anatomy under conflicting notes. These are exploratory constrained-choice findings, not a clinical model ranking.')
    d.pic('assets/prompt_tradeoff.png',45,140,870,335)

    d.new('Keep the decomposition; narrow what it proves','Discussion notes §4 · T01 usable information / PVI', 'For deterministic Z=g(I), fixed question and finite Y under a specified distribution, the four-term identity is valid. Cross entropy equals H(Y|Z)+KL, then chain rules provide the decomposition. The last term includes approximation, fitting and calibration. A fixed-model ablation does not estimate these terms. A failed probe can reflect finite data, pooling or limited capacity; a successful probe can exploit correlations.')
    d.text('Prediction loss = intrinsic + missing context\n                     + lost encoding + predictor mismatch',45,153,870,110,24,TEAL,True)
    d.bullets(['Population identity under explicit assumptions; not a per-case uncertainty label.', 'Frozen-Z Fano bounds do not establish a universal pretraining/post-training boundary.', 'Ablation likelihood gain is model dependence, not Shannon MI or expert evidence sufficiency.'],x=45,y=295,w=865,h=175,size=20)

    d.new('A scoped experiment can establish the causal claim','Proposed study · not an experiment already run', 'Start in a single modality and use source-grounded claims. Prior studies and biopsy reports can be withheld or supplied without editing anatomy. Expert annotation should be blinded and include unavailable, visible, contextual, derived and ambiguous labels. For image changes, use matched cases or validated edits plus irrelevant-region controls. Do not automatically interpret lesion removal as proof of a unique negative diagnosis.')
    d.bullets(['Annotate claims and required evidence on independent, patient/article-separated cases.', 'Manipulate evidence availability: prior study, pathology report, missing panel. Keep questions fixed.', 'Cross supervision alignment with a strong grounding objective at matched data and compute.', 'Measure answerable-case accuracy, unsupported claims, correct updating and irrelevant-change stability together.'],size=21)

    d.new('The primary outcome must require useful answers','Proposed joint evidential success', 'A correct final answer can still contain false premises; evaluate that explicitly. Distinguish withholding evidence from a negative study. If the diagnosis is ambiguous, use acceptable sets or probability labels. Report the four components separately as well as their intersection, clustered by independent patient/article. Pre-register prompts, primary loss, sample-size justification and exclusions before inspecting new model outputs.')
    for i,(t,b) in enumerate([('Sufficient evidence','Answer adequately'),('Irrelevant change','Preserve supported claims'),('Required evidence absent','Qualify / request evidence'),('Valid evidence arrives','Update appropriately')]):
        x=45+(i%2)*445;y=150+(i//2)*160;d.rect(x,y,425,140,'#e5eff1');d.text(t,x+17,y+15,391,45,21,TEAL,True);d.text(b,x+17,y+73,391,55,21)

    d.new('What would weaken the hypothesis?', 'Pre-specified falsification and go/no-go criteria', 'A method should not pass by refusing everything. Compare aligned cleaning with equal-size random cleaning, standard SFT/DPO and relevant MCR or counterfactual controls. Include patient/article clustering and an external-site evaluation. Representation tests become meaningful on independently annotated lesions rather than only broad head/chest anatomy. Conformal calibration comes after score and loss design, with separate calibration patients.')
    d.bullets(['Independent experts find little unavailable-evidence supervision in uniform samples.', 'Aligned supervision fails to beat equal-size random cleaning or existing grounding objectives.', 'Unsupported claims fall only because correct answerable cases are refused.', 'A stronger frozen readout or better crop resolves the issue: investigate access/use before blaming irreversible encoding loss.'],size=21)

    d.new('Advisor discussion: approve the testable core','Full evidence and reproducible audit: advisor_evidence_brief.html', 'Suggested discussion: choose one modality and clinical annotation partner; define source/evidence labels; agree on a budget-matched 2×2 supervision/objective study and joint outcome. Keep causal dominance, coverage deficits and deployment guarantees open. The project is ready to discuss as a hypothesis supported by traceable examples, not as an already established root-cause finding.')
    d.text('Proceed with an evidence-aligned supervision study.',45,157,865,80,30,TEAL,True)
    d.bullets(['Choose one modality and two evidence requirements.', 'Secure independent clinical annotation and provenance checks.', 'Agree on a matched causal ablation and a joint utility/grounding outcome.'],x=45,y=270,w=865,h=170,size=22)
    d.text('The broad framework is reasonable. The decisive contribution still needs an experiment.',45,450,865,32,17,MUTED)

    for title,ids in [('Sources: backbones and direct supervision precedents',['P01','P02','P03','P04','P05','P06','P07']),('Sources: reasoning and benchmark evidence',['P08','P09','P10','P11','P12','P13','P14']),('Sources: close methods',['P15','P16','P17','P18','P19','P20']),('Sources: attribution, abstention and calibration',['P21','P22','P23','T01','T02','T03'])]:
        d.new(title,'Primary-source links; full locators and limitations in paper_evidence.csv','These links identify the primary sources used in the evidence brief. P24 MedBench and P25 the wound study are discussed in the brief with limitations. No cross-paper score ranking is intended.')
        size=14 if len(ids)>8 else 16;h=328/len(ids)
        for i,pid in enumerate(ids):
            p=next(p for p in papers if p['id']==pid);d.text(pid+'  '+p['paper']+'\n'+p['url'],45,146+i*h,870,h-2,size)
    d.save()

def main():
    make_figures();write_tables();make_html();make_slides()
    print('Built HTML evidence brief, 28-source CSV, 11-case manifests, figures and 21-slide PPTX/PDF.')

if __name__=='__main__':main()
